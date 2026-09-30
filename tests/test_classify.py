import os
import tempfile
import unittest

from intouch_analyzer import classify
from intouch_analyzer.analyzer import analyze
from intouch_analyzer.classify import (DEFAULT_RULES, engineering_summary, load_overrides,
                                       load_rules, save_settings, validate_rules)
from intouch_analyzer.export import build_tables, details_rows, report_tsv, summary_rows

DBDUMP = """:IOAccess,Application,Topic
"PLC1","DASABCIP","Line1"
"PLC1_ALT","DASABCIP","Line1"
"SPARE","DASMBTCP","Spare"
:IODisc,Group,Comment,Logged,AlarmState,AccessName,ItemName,ReadOnly
"P1_Run","$System","Pump 1 running",No,On,"PLC1","B3:0/0",No
"P1_Start_Cmd","$System","",No,None,"PLC1","B3:0/1",No
"P1_Reset","$System","",No,None,"PLC1_ALT","B3:0/2",No
"P1_Pulse","$System","Totaliser pulse",No,None,"PLC1","B3:0/3",No
"Mode_Auto","$System","",No,None,"PLC1","B3:0/4",No
:IOReal,Group,Comment,Logged,LogDeadband,MinEU,MaxEU,EngUnits,HiAlarmState,HiHiAlarmState,AccessName,ItemName,ReadOnly
"FT101","$System","Inlet flow",Yes,2.5,0,100,m3/h,On,On,"PLC1","F8:0",Yes
"FT102","$System","Outlet flow",Yes,2.5,0,100,m3/h,None,None,"PLC1","F8:1",Yes
"LT101","$System","Tank level",Yes,0.25,0,5,m,None,None,"PLC1","F8:2",Yes
"Dose_SP","$System","Dosing setpoint",No,0,0,10,mg/l,None,None,"PLC1","F8:3",No
"Speed_Ref","$System","",No,0,0,50,Hz,None,None,"PLC1","F8:4",No
:MemoryReal,Group,Comment,Logged,LogDeadband
"Calc","$System","",No,0
"""

XML = """<Windows>
  <Window Name="Main" WindowType="Replace"><Expr>P1_Run FT101 LT101</Expr></Window>
  <Window Name="Nav" WindowType="Overlay"><Expr>Mode_Auto</Expr></Window>
  <Window Name="Pump_Faceplate"><Script>P1_Start_Cmd = 1; Speed_Ref = 25;
     IF Mode_Auto == 1 THEN Calc = 2; ENDIF;</Script></Window>
  <Window Name="Overview"><Expr>FT102</Expr></Window>
  <Window Name="Template_Old"><Expr>Dose_SP</Expr></Window>
</Windows>
"""


class ClassifyTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.app = os.path.join(self.tmp.name, "app")
        os.makedirs(self.app)
        with open(os.path.join(self.app, "db.csv"), "w") as fh:
            fh.write(DBDUMP)
        with open(os.path.join(self.app, "windows.xml"), "w") as fh:
            fh.write(XML)
        self.res = analyze(self.app)
        self.rules = dict(DEFAULT_RULES)

    def tearDown(self):
        self.tmp.cleanup()

    def test_mimic_types(self):
        eng = engineering_summary(self.res, self.rules, {})
        types = {k.split("|")[1]: v for k, v in eng.mimic_types.items()}
        self.assertEqual(types["Main"], ("Process", "Window type"))
        self.assertEqual(types["Nav"], ("Overlay", "Window type"))
        self.assertEqual(types["Pump_Faceplate"], ("Popup", "Name rule"))
        self.assertEqual(types["Overview"], ("Process", "Default"))
        self.assertEqual(types["Template_Old"], ("Other", "Name rule"))
        eng = engineering_summary(self.res, self.rules,
                                  {"windows.xml|Overview": "Popup"})
        self.assertEqual(eng.mimic_types["windows.xml|Overview"], ("Popup", "Manual"))

    def test_io_classes(self):
        eng = engineering_summary(self.res, self.rules, {})
        c = eng.io_classes
        self.assertEqual(c["p1_run"], "DI")          # input rule
        self.assertEqual(c["p1_start_cmd"], "DO")    # output rule
        self.assertEqual(c["p1_reset"], "DO")
        self.assertEqual(c["p1_pulse"], "PULSE")
        self.assertEqual(c["mode_auto"], "DI")       # compared (==), not written
        self.assertEqual(c["ft101"], "AI")           # ReadOnly
        self.assertEqual(c["dose_sp"], "AO")         # output rule
        self.assertEqual(c["speed_ref"], "AO")       # written in a script
        rules = dict(self.rules, script_writes_are_outputs=False)
        self.assertEqual(engineering_summary(self.res, rules, {}).io_classes["speed_ref"], "AI")

    def test_summary_rows(self):
        eng = engineering_summary(self.res, self.rules, {})
        self.assertEqual(eng.rows(len(self.res.windows), len(self.res.io_tags)), [
            ("Number of mimics", 5),
            ("Mimic Breakdown by type (Process, Popups)",
             "Process - 2 / Popups (On Top + Overlay) - 2"),
            ("Number of data points (I/O tags)", 10),
            ("Breakdown by type (DI / DO / AI / AO)", "Analogues - 5 / Digital - 5"),
            ("   of which (DI / DO / AI / AO)", "AI 3 / AO 2 / DI 2 / DO 2 / PULSE 1"),
            ("Number of configured alarms", 3),
            ("Number of historised points", 3),
            ("Historisation rate", "Model 1-Flow-Change-2.5%-2\nModel 2-Level-Change-5%-1"),
            # PLC1 and PLC1_ALT share application|topic -> one connection; SPARE unused
            ("Number of PLC connections", 1),
        ])
        details = dict(eng.detail_rows())
        self.assertEqual(details["Popup (On Top) mimics"], 1)
        self.assertEqual(details["Overlay mimics"], 1)
        self.assertEqual(details["Other mimics (not counted in the breakdown)"], 1)
        self.assertEqual(details["Tags with alarms"], 2)
        self.assertEqual(details["PLC connections"], "PLC1")

    def test_tables(self):
        eng = engineering_summary(self.res, self.rules, {})
        tables = build_tables(self.res, eng)
        self.assertEqual(tables["Summary"], ([], summary_rows(self.res, eng)))
        self.assertIn(("Tags with alarms", 2), tables["Details"][1])
        self.assertIn(("Total tags", "11"), details_rows(self.res, eng))
        headers, rows = tables["History Models"]
        self.assertEqual(rows[0], ("Model 1", "Flow", "Change-2.5%", 2, "FT101; FT102"))
        headers, rows = tables["Windows"]
        main = next(r for r in rows if r[0] == "Main")
        self.assertEqual(main[headers.index("Mimic Type")], "Process")

    def test_report_tsv(self):
        eng = engineering_summary(self.res, self.rules, {})
        tsv = report_tsv(self.res, eng)
        self.assertTrue(tsv.startswith("Number of mimics\t5\n"))
        self.assertIn('Historisation rate\t"Model 1-Flow-Change-2.5%-2\n'
                      'Model 2-Level-Change-5%-1"\n', tsv)

    def test_excel_report_sheet(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl not installed")
        from intouch_analyzer.export import export_excel
        path = os.path.join(self.tmp.name, "r.xlsx")
        export_excel(self.res, path, engineering_summary(self.res, self.rules, {}))
        wb = openpyxl.load_workbook(path)
        self.assertEqual(wb.sheetnames[:2], ["Summary", "Details"])
        ws = wb["Summary"]
        self.assertEqual(ws.max_row, 9)
        self.assertEqual(ws["A1"].value, "Number of mimics")
        self.assertEqual(ws["B1"].value, 5)  # a number, not text
        self.assertTrue(ws["B8"].alignment.wrap_text)
        self.assertEqual(ws["A1"].border.left.style, "thin")


class SettingsTest(unittest.TestCase):
    def test_round_trip(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "sub", "settings.json")
            rules = dict(DEFAULT_RULES, popup_windows="(?i)^pp_")
            save_settings(rules=rules, path=path)
            save_settings(app_path=tmp, overrides={"a.win|a": "Popup"}, path=path)
            self.assertEqual(load_rules(path)["popup_windows"], "(?i)^pp_")
            self.assertEqual(load_rules(path)["output_tags"], DEFAULT_RULES["output_tags"])
            self.assertEqual(load_overrides(tmp, path), {"a.win|a": "Popup"})
            self.assertEqual(load_overrides("/other", path), {})

    def test_validate(self):
        self.assertEqual(validate_rules(dict(DEFAULT_RULES)), [])
        bad = dict(DEFAULT_RULES, input_tags="(unclosed",
                   measurement_categories=[["X", "[bad"]])
        self.assertEqual(len(validate_rules(bad)), 2)


if __name__ == "__main__":
    unittest.main()
