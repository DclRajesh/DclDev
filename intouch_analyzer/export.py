"""Report tables and export to Excel (openpyxl) or CSV."""

from __future__ import annotations

import csv
import os
from typing import Dict, List, Optional, Sequence, Tuple

from .classify import (EngineeringSummary, alarm_count, engineering_summary, history_model,
                       load_overrides, load_rules, measurement_category, override_key)
from .history import history_sort_key
from .models import AnalysisResult

Table = Tuple[List[str], List[Sequence]]


def _yn(flag: bool) -> str:
    return "Yes" if flag else "No"


def default_engineering(r: AnalysisResult) -> EngineeringSummary:
    return engineering_summary(r, load_rules(), load_overrides(r.app_path))


def summary_rows(r: AnalysisResult,
                 eng: Optional[EngineeringSummary] = None) -> List[Tuple[str, str]]:
    """Engineering summary (mimics, DI/DO/AI/AO, alarms, history, PLCs)
    followed by detailed statistics. Values may contain newlines."""
    eng = eng or default_engineering(r)
    return (eng.rows(len(r.windows), len(r.io_tags))
            + [("", ""), ("Detailed statistics", "")] + detail_rows(r))


def detail_rows(r: AnalysisResult) -> List[Tuple[str, str]]:
    per_access = r.tags_per_access_name()
    rows = [
        ("Application folder", r.app_path),
        ("Tag database source", r.tag_db_source),
        ("Total tags", str(len(r.tags))),
        ("I/O tags", str(len(r.io_tags))),
        ("Access names", str(len(r.access_names))),
        ("Windows (mimics)", str(len(r.windows))),
        ("Other files referencing tags", str(len(r.other_sources))),
        ("Tags used in windows/scripts", str(len(r.tags) - len(r.unused_tags()))),
        ("Unused tags", str(len(r.unused_tags()))),
        ("Duplicate I/O addresses", str(len(r.duplicate_io_addresses()))),
        ("Alarmed tags", str(sum(1 for t in r.tags.values() if t.has_alarm))),
        ("Historised (logged) tags", str(sum(1 for t in r.tags.values() if t.is_logged))),
        ("History source", r.history_source or "none"),
    ]
    for desc, total, detail in history_rows(r):
        rows.append((f"  History: {desc}",
                     f"{total} tag{'' if total == 1 else 's'} ({detail})"))
    for ttype, n in sorted(r.type_counts().items()):
        rows.append((f"  Tag type: {ttype}", str(n)))
    for acc in sorted(r.access_names.values(), key=lambda a: a.key):
        rows.append((f"  Access name: {acc.name}", f"{per_access.get(acc.key, 0)} tags"))
    return rows


def history_rows(r: AnalysisResult) -> List[Tuple[str, int, str]]:
    """[(history description, tag count, 'I/O a, Memory b')] fastest rate first."""
    counts = r.history_counts()
    return [(d, sum(counts[d].values()),
             ", ".join(f"{c} {n}" for c, n in sorted(counts[d].items())))
            for d in sorted(counts, key=history_sort_key)]


def build_tables(r: AnalysisResult,
                 eng: Optional[EngineeringSummary] = None) -> Dict[str, Table]:
    eng = eng or default_engineering(r)
    rules = load_rules()
    usage = r.tag_usage()
    win_usage = r.window_usage()
    per_access = r.tags_per_access_name()
    tables: Dict[str, Table] = {}

    tables["Summary"] = (["Item", "Value"], summary_rows(r, eng))

    tag_rows = []
    for t in sorted(r.tags.values(), key=lambda t: t.key):
        wins = win_usage.get(t.key, [])
        others = [s.name for s in usage.get(t.key, []) if not s.is_window]
        tag_rows.append((t.name, t.tag_type, t.category, t.access_name, t.item_name,
                         t.group, t.comment, _yn(t.has_alarm), t.history or "Not logged",
                         _yn(t.key in usage), len(wins), "; ".join(sorted(wins)),
                         "; ".join(sorted(others)), eng.io_classes.get(t.key, ""),
                         alarm_count(t),
                         measurement_category(t, rules) if t.is_logged else "",
                         history_model(t) if t.is_logged else ""))
    tables["Tags"] = (["Tag", "Type", "Category", "Access Name", "Item", "Group",
                       "Comment", "Alarm", "History", "Used", "# Windows",
                       "Windows", "Other Files", "I/O Class", "Alarm Conditions",
                       "Measurement", "History Model"], tag_rows)

    tables["Access Names"] = (
        ["Access Name", "Application", "Topic", "Advise Active", "Protocol",
         "Secondary Application", "Secondary Topic", "# Tags"],
        [(a.name, a.application, a.topic, a.advise_active, a.protocol,
          a.sec_application, a.sec_topic, per_access.get(a.key, 0))
         for a in sorted(r.access_names.values(), key=lambda a: a.key)])

    counts = r.history_counts()
    categories = sorted({c for cats in counts.values() for c in cats})
    tables["History Rates"] = (
        ["History / Storage Rate", "Tags"] + categories,
        [(d, sum(counts[d].values())) + tuple(counts[d].get(c, 0) for c in categories)
         for d in sorted(counts, key=history_sort_key)])

    model_tags: Dict[Tuple[str, str], List[str]] = {}
    for t in r.tags.values():
        if t.is_logged:
            model_tags.setdefault((measurement_category(t, rules), history_model(t)),
                                  []).append(t.name)
    tables["History Models"] = (
        ["Model", "Measurement", "History", "Points", "Tags"],
        [(f"Model {i}", cat, model, n, "; ".join(sorted(model_tags.get((cat, model), []),
                                                       key=str.lower)))
         for i, (cat, model, n) in enumerate(eng.history_models, start=1)])

    win_rows = []
    for w in sorted(r.windows, key=lambda w: w.name.lower()):
        io = [k for k in w.tag_refs if k in r.tags and r.tags[k].is_io]
        accesses = sorted({r.tags[k].access_name for k in io if r.tags[k].access_name})
        win_rows.append((w.name, w.kind, w.path, len(w.tag_refs), len(io),
                         "; ".join(accesses), "; ".join(sorted(w.opens_windows)),
                         "; ".join(sorted(w.remote_refs)), len(w.unresolved))
                        + eng.mimic_types.get(override_key(w), ("", "")))
    tables["Windows"] = (["Window", "Source", "File", "# Tags", "# I/O Tags",
                          "Access Names", "Opens Windows", "Remote References",
                          "# Undefined Refs", "Mimic Type", "Type From"], win_rows)

    wt_rows = []
    for w in sorted(r.sources, key=lambda s: (not s.is_window, s.name.lower())):
        for key, count in sorted(w.tag_refs.items()):
            t = r.tags.get(key)
            wt_rows.append((w.name, w.kind, t.name if t else key,
                            t.tag_type if t else "", t.access_name if t else "",
                            t.item_name if t else "", count))
    tables["Window-Tag Map"] = (["Window/File", "Source", "Tag", "Type",
                                 "Access Name", "Item", "References"], wt_rows)

    dup_rows = [(addr, t.name, t.tag_type)
                for addr, dup in r.duplicate_io_addresses() for t in dup]
    tables["Duplicate IO"] = (["I/O Address", "Tag", "Type"], dup_rows)

    tables["Issues"] = (["Severity", "Category", "Item", "Detail"],
                        [(i.severity, i.category, i.item, i.detail) for i in r.issues()])
    return tables


def export_excel(r: AnalysisResult, path: str,
                 eng: Optional[EngineeringSummary] = None) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise RuntimeError("Excel export needs openpyxl: pip install openpyxl") from exc

    wb = Workbook()
    wb.remove(wb.active)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    wrap = Alignment(wrap_text=True, vertical="top")
    for name, (headers, rows) in build_tables(r, eng).items():
        ws = wb.create_sheet(name[:31])
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
        for row in rows:
            ws.append(list(row))
            if any(isinstance(v, str) and "\n" in v for v in row):
                for cell in ws[ws.max_row]:
                    cell.alignment = wrap
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for idx, h in enumerate(headers, start=1):
            width = max([len(str(h))] + [len(line) for row in rows[:500]
                                          for line in str(row[idx - 1]).splitlines() or [""]])
            ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 8), 60)
    wb.save(path)


def export_csv(r: AnalysisResult, folder: str,
               eng: Optional[EngineeringSummary] = None) -> List[str]:
    os.makedirs(folder, exist_ok=True)
    written = []
    for name, (headers, rows) in build_tables(r, eng).items():
        path = os.path.join(folder, name.replace(" ", "_").lower() + ".csv")
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh)
            w.writerow(headers)
            w.writerows(rows)
        written.append(path)
    return written
